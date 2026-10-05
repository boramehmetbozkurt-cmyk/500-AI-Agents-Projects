from __future__ import annotations

from io import BytesIO
import re
from typing import Any, TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

if TYPE_CHECKING:
    from app.unified_answer import UnifiedAnswerResponse

_EXCEL_FORMULA_PREFIXES = ("=", "+", "-", "@")
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_MAX_CELL_CHARS = 32767

_HEADER_FILL = PatternFill("solid", fgColor="17365D")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
_TITLE_FONT = Font(color="17365D", bold=True, size=14)


def _safe_cell(value: Any) -> Any:
    """Return a value safe for an XLSX cell, including CSV/Excel formula-injection defense."""
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = _ILLEGAL_XML.sub("", str(value))[:_MAX_CELL_CHARS]
    if text.startswith(_EXCEL_FORMULA_PREFIXES):
        text = "'" + text
    return text


def _join(values: Any) -> str:
    if not values:
        return ""
    if isinstance(values, (list, tuple, set)):
        return ", ".join(str(v) for v in values)
    return str(values)


def _write_table(ws: Worksheet, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append([_safe_cell(x) for x in headers])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_safe_cell(x) for x in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _fit_columns(ws)


def _fit_columns(ws: Worksheet) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            line_width = max((len(line) for line in value.splitlines()), default=0)
            widths[cell.column] = min(max(widths.get(cell.column, 0), line_width + 2), 70)
    for index, width in widths.items():
        ws.column_dimensions[get_column_letter(index)].width = max(10, width)


def _summary_sheet(wb: Workbook, result: "UnifiedAnswerResponse") -> None:
    ws = wb.active
    ws.title = "Ozet"
    ws.merge_cells("A1:D1")
    ws["A1"] = "ORBYTHRA Arama ve Claim Disa Aktarimi"
    ws["A1"].fill = _TITLE_FILL
    ws["A1"].font = _TITLE_FONT
    ws["A1"].alignment = Alignment(vertical="center")

    receipt = result.receipt or {}
    verification = result.verification or {}
    rows = [
        ["Sorgu", result.query],
        ["Durum", result.status],
        ["Cevap", result.answer or ""],
        ["Guven", result.confidence],
        ["Etkin mod", result.effective_mode or ""],
        ["Kaynak sayisi", len(result.sources)],
        ["Claim sayisi", len(result.claims)],
        ["Dogrulama gecti", verification.get("passed")],
        ["Model assurance", verification.get("model_assurance_status")],
        ["V1 disposition", verification.get("v1_disposition")],
        ["Kanıt kapisi gecis orani", verification.get("evidence_gate_pass_rate")],
        ["Basarisiz kapilar", _join(verification.get("failed_gates"))],
        ["Request ID", receipt.get("request_id")],
        ["Release SHA", receipt.get("release_sha")],
        ["Configuration SHA-256", receipt.get("configuration_sha256")],
        ["Output SHA-256", receipt.get("output_sha256")],
        ["Evidence manifest SHA-256", receipt.get("evidence_manifest_sha256")],
        ["Secilen araclar", _join(receipt.get("selected_tools"))],
        ["Uyarilar", _join(result.warnings)],
        ["Sinir", result.boundary],
    ]
    ws.append([])
    ws.append(["Alan", "Deger"])
    for cell in ws[3]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row in rows:
        ws.append([_safe_cell(row[0]), _safe_cell(row[1])])
    for row in ws.iter_rows(min_row=4, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A4"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95


def build_unified_answer_xlsx(result: "UnifiedAnswerResponse") -> bytes:
    wb = Workbook()
    _summary_sheet(wb, result)

    sources = wb.create_sheet("Arama Sonuclari")
    _write_table(
        sources,
        [
            "No", "Source ID", "Baslik", "Yayinci", "URL", "Kaynak Tipi",
            "Dogrudan Kanit", "Kanit Durumu", "Yayin Tarihi", "Alinma Zamani (Unix)", "Snippet",
        ],
        [
            [
                index,
                source.source_id,
                source.title,
                source.publisher or "",
                source.url,
                source.source_type,
                source.direct_fetch,
                getattr(source, "evidence_status", "DIRECT_EVIDENCE" if source.direct_fetch else "DISCOVERY_ONLY"),
                source.published_at or "",
                source.retrieved_at,
                source.snippet or "",
            ]
            for index, source in enumerate(result.sources, start=1)
        ],
    )

    claims = wb.create_sheet("Claimler")
    claim_rows: list[list[Any]] = []
    for index, claim in enumerate(result.claims, start=1):
        claim_rows.append([
            index,
            claim.get("claim", ""),
            claim.get("status", ""),
            claim.get("nli_label", ""),
            claim.get("nli_entailment"),
            claim.get("nli_contradiction"),
            _join(claim.get("supporting_sources")),
            _join(claim.get("contradicting_sources")),
            claim.get("nli_verifier", ""),
        ])
    _write_table(
        claims,
        [
            "No", "Claim", "Durum", "NLI Label", "NLI Entailment",
            "NLI Contradiction", "Destekleyen Source ID", "Celisen Source ID", "NLI Verifier",
        ],
        claim_rows,
    )
    claims.column_dimensions["B"].width = 80
    claims.column_dimensions["G"].width = 38
    claims.column_dimensions["H"].width = 38

    verification = wb.create_sheet("Dogrulama")
    failed = result.verification.get("failed_gates") or []
    verification_rows = [
        ["passed", result.verification.get("passed")],
        ["model_assurance_status", result.verification.get("model_assurance_status")],
        ["v1_disposition", result.verification.get("v1_disposition")],
        ["direct_source_count", result.verification.get("direct_source_count")],
        ["evidence_gate_pass_rate", result.verification.get("evidence_gate_pass_rate")],
        ["failed_gates", _join(failed)],
        ["warnings", _join(result.warnings)],
    ]
    _write_table(verification, ["Alan", "Deger"], verification_rows)

    pipeline = wb.create_sheet("Pipeline")
    _write_table(
        pipeline,
        ["No", "Adim", "Durum", "Detay", "Endpoint"],
        [[i, step.step, step.status, step.detail, step.endpoint or ""] for i, step in enumerate(result.pipeline, start=1)],
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
