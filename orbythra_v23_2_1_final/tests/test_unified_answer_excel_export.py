from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.unified_answer_export import build_unified_answer_xlsx


def sample_result() -> SimpleNamespace:
    source = SimpleNamespace(
        source_id="source-1",
        title='=HYPERLINK("https://evil.example","bad")',
        url="https://example.com/source",
        publisher="example.com",
        published_at="2026-10-05",
        retrieved_at=1_800_000_000,
        snippet="A source-backed fact.",
        source_type="test",
        direct_fetch=True,
    )
    step = SimpleNamespace(step="searching_sources", status="PASS", detail="1 source found", endpoint="/search")
    return SimpleNamespace(
        query="ORBYTHRA test sorgusu",
        status="PASS",
        answer="Kaynakla doğrulanmış test cevabı.",
        confidence=0.93,
        effective_mode="WEB_RESEARCH",
        sources=[source],
        claims=[{
            "claim": "+2+3",
            "status": "SUPPORTED",
            "nli_label": "ENTAILMENT",
            "nli_entailment": 0.97,
            "nli_contradiction": 0.01,
            "supporting_sources": ["source-1"],
            "contradicting_sources": [],
            "nli_verifier": "test-verifier",
        }],
        verification={
            "passed": True,
            "model_assurance_status": "PASS",
            "v1_disposition": "ANSWER",
            "direct_source_count": 1,
            "evidence_gate_pass_rate": 1.0,
            "failed_gates": [],
        },
        pipeline=[step],
        receipt={
            "request_id": "req-export-001",
            "release_sha": "1" * 40,
            "configuration_sha256": "2" * 64,
            "output_sha256": "3" * 64,
            "evidence_manifest_sha256": "4" * 64,
            "selected_tools": ["web_search"],
        },
        warnings=[],
        boundary="Scoped verification boundary.",
    )


def test_excel_export_builds_real_xlsx_with_sources_and_claims() -> None:
    payload = build_unified_answer_xlsx(sample_result())
    assert payload[:2] == b"PK"

    wb = load_workbook(BytesIO(payload), data_only=False)
    assert wb.sheetnames == ["Ozet", "Arama Sonuclari", "Claimler", "Dogrulama", "Pipeline"]
    assert wb["Arama Sonuclari"]["B2"].value == "source-1"
    assert wb["Arama Sonuclari"]["C2"].value.startswith("'=")
    assert wb["Claimler"]["B2"].value.startswith("'+")
    assert wb["Claimler"]["G2"].value == "source-1"
    assert wb["Ozet"]["B4"].value == "ORBYTHRA test sorgusu"


def test_excel_export_contract_is_authenticated_and_downloadable() -> None:
    source = open("app/unified_answer_policy.py", encoding="utf-8").read()
    assert '@router.post("/export/xlsx")' in source
    assert "Depends(_REQUIRE_OPERATOR)" in source
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in source
    assert "Content-Disposition" in source
